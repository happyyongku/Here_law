from openai import OpenAI
from fpdf import FPDF
from docx import Document
from datetime import datetime

# OpenAI 클라이언트 생성 (API 키 필요 시 api_key 파라미터를 추가하세요)
client = OpenAI()  # 여기에 자신의 OpenAI API 키를 입력하세요.

def gather_user_info():
    print("안녕하세요! 소장 작성을 위해 몇 가지 질문을 드리겠습니다.\n")

    case_title = input("사건명을 입력해주세요: ")
    plaintiff = input("원고의 이름을 입력해주세요: ")
    plaintiff_address = input("원고의 주소를 입력해주세요: ")
    plaintiff_phone = input("원고의 전화번호를 입력해주세요: ")
    defendant = input("피고의 이름을 입력해주세요: ")
    defendant_address = input("피고의 주소를 입력해주세요: ")
    defendant_phone = input("피고의 전화번호를 입력해주세요: ")
    court_name = input("관할 법원의 이름을 입력해주세요: ")
    case_details = input("사건의 상세 내용을 입력해주세요: ")

    return {
        "case_title": case_title,
        "plaintiff": plaintiff,
        "plaintiff_address": plaintiff_address,
        "plaintiff_phone": plaintiff_phone,
        "defendant": defendant,
        "defendant_address": defendant_address,
        "defendant_phone": defendant_phone,
        "court_name": court_name,
        "case_details": case_details,
    }

def generate_content(user_info):
    prompt = f"""
다음 정보를 바탕으로 한국어 소장에 필요한 청구 취지, 청구 원인, 첨부 서류를 작성해주세요.

사건명: {user_info['case_title']}
원고: {user_info['plaintiff']} (주소: {user_info['plaintiff_address']}, 전화번호: {user_info['plaintiff_phone']})
피고: {user_info['defendant']} (주소: {user_info['defendant_address']}, 전화번호: {user_info['defendant_phone']})
사건 내용: {user_info['case_details']}
소장은 정식 법률 문서 형식으로 작성되어야 하며, 필요한 모든 법적 요소를 포함해야 합니다.
[청구 취지], [청구 원인], [첨부 서류] 이런식으로 작성해 주세요
청구 취지는 1. 피고는 원고에게 ...하라, 2. 소송비용은 피고의 부담으로 한다. 3. 제 1항은 가집행 할 수 있다. 라는 판결을 구합니다. 이런 형식이어야 합니다
"""

    response = client.chat.completions.create(
        model="gpt-4",
        messages=[
            {"role": "system", "content": "당신은 한국의 법률 문서 작성 전문가입니다."},
            {"role": "user", "content": prompt}
        ]
    )

    generated_text = response.choices[0].message.content
    print(generated_text)

    # 생성된 텍스트를 섹션별로 분리
    sections = {"청구 취지": "", "청구 원인": "", "첨부 서류": ""}
    current_section = None
    for line in generated_text.split('\n'):
        line = line.strip()
        if line.startswith("[청구 취지]") or line.startswith("청구 취지") or line.startswith("청구취지") or line.startswith("청구의 취지"):
            current_section = "청구 취지"
            continue
        elif line.startswith("[청구 원인]") or line.startswith("청구 원인") or line.startswith("청구원인") or line.startswith("청구의 원인"):
            current_section = "청구 원인"
            continue
        elif line.startswith("[첨부 서류]") or line.startswith("첨부 서류") or line.startswith("첨부서류"):
            current_section = "첨부 서류"
            continue
        elif line == "":
            continue
        if current_section:
            sections[current_section] += line + '\n'

    return sections

def save_to_pdf(user_info, generated_content, filename="소장.pdf"):
    pdf = FPDF()
    pdf.add_page()
    pdf.add_font('NanumGothic', '', '/home/qudrb0107/S11P21B109/datas/code/NanumGothic.ttf', uni=True)  # 폰트 파일 필요
    pdf.set_font('NanumGothic', size=12)

    # 소장 제목 (가운데 정렬, 큰 글씨)
    pdf.set_font_size(16)
    pdf.cell(0, 10, txt="소      장", ln=True, align='C')

    pdf.ln(10)  # 줄 간격

    # 원고 정보
    pdf.cell(30, 10, txt="원    고 :", ln=False)
    pdf.cell(0, 10, txt=f"{user_info['plaintiff']} (전화번호: {user_info['plaintiff_phone']})", ln=True)
    pdf.cell(30, 10, txt="주    소 :", ln=False)
    pdf.cell(0, 10, txt=user_info['plaintiff_address'], ln=True)

    # 피고 정보
    pdf.cell(30, 10, txt="피    고 :", ln=False)
    pdf.cell(0, 10, txt=f"{user_info['defendant']} (전화번호: {user_info['defendant_phone']})", ln=True)
    pdf.cell(30, 10, txt="주    소 :", ln=False)
    pdf.cell(0, 10, txt=user_info['defendant_address'], ln=True)
    
    pdf.ln(20)  # 줄 간격
    
    # 사건명
    pdf.set_font_size(12)
    pdf.cell(0, 10, txt=user_info['case_title'], ln=True)


    pdf.ln(20)

    # 청구 취지
    pdf.set_font_size(14)
    pdf.cell(0, 10, txt="청 구  취 지", ln=True, align='C')
    pdf.set_font_size(12)
    pdf.multi_cell(0, 10, txt=generated_content['청구 취지'], align='J')

    pdf.ln(10)

    # 청구 원인
    pdf.set_font_size(14)
    pdf.cell(0, 10, txt="청 구  원 인", ln=True, align='C')
    pdf.set_font_size(12)
    pdf.multi_cell(0, 10, txt=generated_content['청구 원인'], align='J')

    pdf.ln(10)

    # 첨부 서류
    pdf.set_font_size(14)
    pdf.cell(0, 10, txt="첨 부  서 류", ln=True, align='C')
    pdf.set_font_size(12)
    attachments = generated_content['첨부 서류'].split('\n')
    for attachment in attachments:
        if attachment.strip() != '':
            pdf.cell(0, 10, txt=f"- {attachment.strip()}", ln=True)

    pdf.ln(20)

    # 날짜 및 원고 서명 (오른쪽 정렬)
    from datetime import datetime
    today = datetime.today().strftime('%Y    .   %m    .   %d    .')
    pdf.cell(0, 10, txt=today, ln=True, align='R')
    pdf.cell(0, 10, txt="원고           (인)", ln=True, align='R')
    pdf.cell(0, 10, txt=user_info['plaintiff'], ln=True, align='R')

    pdf.ln(10)

    # 법원 이름 (오른쪽 정렬)
    pdf.cell(0, 10, txt=f"{user_info['court_name']} 귀중", ln=True, align='R')

    pdf.output(filename, 'F')
    print(f"\n'{filename}' 파일로 저장되었습니다.")

def main():
    user_info = gather_user_info()
    generated_content = generate_content(user_info)
    print("\n생성된 청구 취지, 청구 원인, 첨부 서류:\n")
    for section, content in generated_content.items():
        print(f"{section}:\n{content}\n")
    save_to_pdf(user_info, generated_content)

if __name__ == "__main__":
    main()
